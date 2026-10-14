"""Produce editable Word report using the reviewed Week 3 report content."""
from pathlib import Path
import importlib.util, re, html
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
TMP=ROOT/'tmp/word-week3'
TMP.mkdir(parents=True,exist_ok=True)
OUT=ROOT/'output/docx/rescuelink-week3-report.docx'
OUT.parent.mkdir(parents=True,exist_ok=True)
spec=importlib.util.spec_from_file_location('week3',ROOT/'tool/create_week3_latest.py')
source=importlib.util.module_from_spec(spec); spec.loader.exec_module(source)
sections=[]
class DummyCanvas:
    def __getattr__(self,key): return lambda *a,**kw: None
class Collector:
    def __init__(self,*args): self.c=DummyCanvas(); self.y=700
    def start(self,title): sections.append({'title':title,'items':[]})
    def p(self,text,style=None,gap=10): sections[-1]['items'].append(('p',text,style))
    def table(self,headers,rows,widths): sections[-1]['items'].append(('table',headers,rows,widths))
    def save(self): pass
source.PDF=Collector
source.report()

def plain(text):
    return html.unescape(re.sub(r'<[^>]+>','',text.replace('<br/>','\n')))
def fonts(run,size=None):
    run.font.name='Leelawadee UI'
    if size: run.font.size=Pt(size)
    props=run._element.get_or_add_rPr()
    rf=props.rFonts
    for attr in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+attr),'Leelawadee UI')
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');lang.set(qn('w:eastAsia'),'th-TH');lang.set(qn('w:bidi'),'th-TH');props.append(lang)
    if size:
        cs=OxmlElement('w:szCs');cs.set(qn('w:val'),str(int(size*2)));props.append(cs)

doc=Document(); sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(17);sec.bottom_margin=Mm(17)
sec.left_margin=Mm(16);sec.right_margin=Mm(16)
sec.header_distance=Mm(7);sec.footer_distance=Mm(7)
for name,size in [('Normal',10.5),('Title',23),('Heading 1',16),('Heading 2',12)]:
    style=doc.styles[name];style.font.name='Leelawadee UI';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().rFonts
    for attr in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+attr),'Leelawadee UI')
    style.paragraph_format.space_after=Pt(8)
    style.paragraph_format.line_spacing=1.12
    style.paragraph_format.widow_control=True
doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(6)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(6)
for style in doc.styles:
    for border in style.element.findall('.//' + qn('w:pBdr')):
        border.getparent().remove(border)

header=sec.header.paragraphs[0]
fonts(header.add_run('RescueLink   Advanced Programming   Week 3'),8)
header.paragraph_format.space_after=Pt(0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fonts(footer.add_run('หน้า '),8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

# Keep the original ER diagram at print resolution; all report text and tables
# are native, editable Word objects.
diagram=TMP/'schema.png'
im=Image.open(TMP/'source-schema-3.png')
im.crop((int(im.width*.068),int(im.height*.158),int(im.width*.90),int(im.height*.497))).save(diagram)

for index,section in enumerate(sections):
    if index==0:
        p=doc.add_paragraph('รายงานความคืบหน้า RescueLink สัปดาห์ที่ 3',style='Title')
        for run in p.runs: fonts(run,23)
        p=doc.add_paragraph('รายวิชา Advanced Programming\nวันที่ 10 ตุลาคม 2569')
        for run in p.runs: fonts(run,10.5)
    heading=re.sub(r'^(\d+)\.\s*',r'\1 ',section['title']).replace(':','').replace(',','').replace('/',' ')
    p=doc.add_paragraph(heading,style='Heading 1')
    if index: p.paragraph_format.page_break_before=True
    for run in p.runs: fonts(run,16)
    if index==2:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(diagram),width=Mm(166))
        p.paragraph_format.space_after=Pt(10)
    for item in section['items']:
        if item[0]=='p':
            text=plain(item[1])
            if text.startswith('PDF นี้เป็นรายงานประกอบส่ง'):
                text=text.replace('PDF นี้เป็นรายงานประกอบส่ง','รายงานฉบับนี้ใช้ประกอบการส่งงาน')
            # The author can fill in their name without a guessed identity.
            if text.startswith('ชื่อผู้จัดทำและรหัสนักศึกษาไม่ได้ปรากฏ'): continue
            is_sub=item[2] is source.SUB
            p=doc.add_paragraph(text,style='Heading 2' if is_sub else 'Normal')
            if item[2] is source.SMALL: p.paragraph_format.space_after=Pt(6)
            for run in p.runs: fonts(run,12 if is_sub else (9.5 if item[2] is source.SMALL else 10.5))
        else:
            _,headers,rows,widths=item
            t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            avail=178
            for col,width in zip(t.columns,widths): col.width=Mm(avail*width/sum(widths))
            props=t._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders)
            margins=OxmlElement('w:tblCellMar')
            for side,value in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            for n,values in enumerate([headers]+rows):
                row=t.rows[0] if n==0 else t.add_row()
                pr=row._tr.get_or_add_trPr()
                no_split=OxmlElement('w:cantSplit');pr.append(no_split)
                if n==0:
                    repeat=OxmlElement('w:tblHeader');pr.append(repeat)
                for i,(cell,value) in enumerate(zip(row.cells,values)):
                    cell.width=Mm(avail*widths[i]/sum(widths));cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF1F9' if n==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(shade)
                    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.05
                    p.paragraph_format.keep_with_next=False
                    run=p.add_run(plain(str(value)));fonts(run,9);run.bold=n==0
            spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(4);spacer.paragraph_format.line_spacing=Pt(4)

doc.core_properties.title='รายงานความคืบหน้า RescueLink สัปดาห์ที่ 3'
doc.core_properties.subject='Advanced Programming Week 3 database schema and CRUD progress'
doc.core_properties.author='RescueLink project'
doc.save(OUT)
print(OUT)
print(f'{len(sections)} sections; {len(doc.tables)} editable tables; {len(doc.inline_shapes)} diagram')
